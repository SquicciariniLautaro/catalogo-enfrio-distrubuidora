"""Arma datos/productos.json a partir de los Excel de precios de EnFrío.

Uso (desde la carpeta del proyecto):
    python scripts/importar_excel.py "CATALOGO NUEVO.xlsx" "CATALOGO A GRANEL.xlsx"

Necesita openpyxl (python -m pip install openpyxl).

Se usó para la carga inicial. Después los precios se manejan desde el panel
(admin.html); volver a correrlo REEMPLAZA todo lo que se haya cambiado ahí.
"""

import json
import re
import sys
from pathlib import Path

import openpyxl

# Si el Excel no trae el descuento para cliente fijo, se usa este.
DESCUENTO_FIJO_POR_DEFECTO = 0.05

ORDEN_CATEGORIAS = [
    "Pescados y mariscos", "Pollo", "Croquetas y bocaditos", "Veggie", "Papas",
    "Pastas", "Tartas", "Panificados", "Sin gluten", "Frutas", "Helados y postres", "Bebidas",
]

MARCAS = [
    ("CASA VEGANA", "Casa Vegana"), ("MAMA SABE", "Mamá Sabe"), ("FINCA BALCARCE", "Finca Balcarce"),
    ("ALL FRONZEN", "All Frozen"), ("ALL FROZEN", "All Frozen"), ("FARM FRITES", "Farm Frites"),
    ("FARMFRITES", "Farm Frites"), ("STAR FRITES", "Farm Frites"), ("ARTICO", "Ártico"),
    ("SUPERBE", "Superbe"), ("KARINAT", "Karinat"), ("KUIBI", "Kubi"), ("KUBI", "Kubi"),
    ("CASERIANA", "Caseriana"), ("FRANUI", "Franui"), ("QUEM", "Qüem"), ("MYDIBEL", "Mydibel"),
]

# Marca según el prefijo del código de barras, para filas que no la nombran.
MARCA_POR_CODIGO = [
    ("7794058", "Ártico"), ("7798146", "Superbe"), ("7798141", "Karinat"),
    ("7798302", "Finca Balcarce"), ("0724751", "Caseriana"), ("7798147", "Franui"),
]

ERRATAS = {
    "PREMUIN": "PREMIUM", "PREMIUN": "PREMIUM", "LAGOSNTINOS": "LANGOSTINOS",
    "LANGOSITNO": "LANGOSTINO", "DEVENADO": "DESVENADO", "ESPINCA": "ESPINACA",
    "BOLONEGSA": "BOLOGNESA", "FRONZEN": "FROZEN", "BOQUE": "BLOQUE", "ALA ROMANA": "A LA ROMANA",
    "PRODCUTO": "PRODUCTO", "MARRACUYA": "MARACUYA", "C/MARACUYA": "MARACUYA CON SEMILLA",
    "S/MARACUYA": "MARACUYA SIN SEMILLA", "MARACUYA C/S": "MARACUYA CON SEMILLA",
    "MARACUYA S/S": "MARACUYA SIN SEMILLA", "C/CASCARA": "CON CASCARA", "CONG.": "CONGELADO",
    "EMP.": "ENVASADO", "DESP.": "DESPINADO", "V/P": "AL VACIO", "VAC PACK": "AL VACIO",
}

# Tildes que no vienen en el Excel.
TILDES = {
    "medallon": "medallón", "jamon": "jamón", "salmon": "salmón", "camaron": "camarón",
    "arandanos": "arándanos", "arandano": "arándano", "limon": "limón", "budin": "budín",
    "noquis": "ñoquis", "maiz": "maíz", "clasico": "clásico", "mediterranea": "mediterránea",
    "hungaro": "húngaro", "lentejon": "lentejón", "patagonico": "patagónico",
    "patagonica": "patagónica", "maracuya": "maracuyá", "brocoli": "brócoli",
    "tentaculo": "tentáculo", "atun": "atún", "pacifico": "pacífico", "sandwich": "sándwich",
    "galleton": "galletón", "sesamo": "sésamo", "baston": "bastón", "vacio": "vacío",
    "cascara": "cáscara", "estandar": "estándar", "porcion": "porción", "bolognesa": "boloñesa",
    "lasagna": "lasaña",
}

MINUSCULAS = {"de", "del", "y", "con", "a", "la", "al", "en", "x", "sin", "e", "o", "el", "los", "las", "tipo"}
SIGLAS = {"IQF", "TACC", "ARG", "BK-S", "VR-S", "PF", "C/S", "S/S", "C/", "S/", "V/P", "MM", "BL"}


def texto(valor):
    return "" if valor is None else str(valor).strip()


def numero(valor):
    if isinstance(valor, (int, float)) and valor > 0:
        return valor
    return None


def corregir(nombre):
    nombre = " ".join(nombre.upper().split())
    for mal, bien in ERRATAS.items():
        nombre = re.sub(rf"\b{re.escape(mal)}\b", bien, nombre)
    return nombre


def detectar_marca(nombre, codigo, por_defecto=""):
    for clave, marca in MARCAS:
        if clave in nombre:
            return marca
    for prefijo, marca in MARCA_POR_CODIGO:
        if codigo.lstrip("0").startswith(prefijo.lstrip("0")):
            return marca
    return por_defecto


def quitar_marca(nombre):
    for clave, _ in MARCAS:
        nombre = re.sub(rf"\b{clave}\b\s*[-:]?\s*", "", nombre)
    return nombre.strip(" -:")


def titulo(nombre):
    palabras = []
    for i, palabra in enumerate(nombre.split()):
        if palabra in SIGLAS or re.fullmatch(r"\d+([.,/-]\d+)*", palabra):
            palabras.append(palabra)
        elif i > 0 and palabra.lower() in MINUSCULAS:
            palabras.append(palabra.lower())
        else:
            minuscula = TILDES.get(palabra.lower(), palabra.lower())
            palabras.append(minuscula[0].upper() + minuscula[1:])
    resultado = " ".join(palabras)
    resultado = re.sub(r"(\d+)\s*Grs?\b", r"\1 g", resultado)
    resultado = re.sub(r"(\d+)\s*Kg\b", r"\1 kg", resultado)
    resultado = re.sub(r" x Unidad\b", " c/u", resultado)
    return " ".join(resultado.replace(" - ", " ").split())


def peso(cantidad, unidad):
    cantidad = cantidad.replace(",", ".")
    if unidad in ("K", "KG", "KGS"):
        return f"{cantidad.rstrip('0').rstrip('.') if '.' in cantidad else cantidad} kg"
    if unidad in ("LTS", "LT", "L"):
        return f"{cantidad} L"
    return f"{cantidad} g"


def separar_presentacion(nombre):
    """Saca peso, unidades y bulto del nombre. Devuelve (nombre, presentacion, unidades_bulto)."""
    partes = []
    bulto = ""
    m = re.search(r"\((\d+)\)", nombre)
    if m:
        bulto = m.group(1)
        nombre = (nombre[:m.start()] + " " + nombre[m.end():]).strip()
    nombre = nombre.replace("*", " ").replace('"', " ")
    nombre = re.sub(r"\bX(\d{3})\b", r"\1G", nombre)

    m = re.search(r"\b(\d+)\s*U(?:ND|NIDADES)?\b\.?|\bX\s?(\d{1,2})\b(?!\s*(?:G|GR|KG|K)\b)|\b(DOCENA)\b", nombre)
    if m:
        cantidad = m.group(1) or m.group(2)
        partes.append("12 u." if m.group(3) else f"{cantidad} u.")
        nombre = (nombre[:m.start()] + nombre[m.end():])

    pesos = list(re.finditer(r"(?:X\s*)?(\d+(?:[.,]\d+)?)\s*(KGS?|K|GRS?|G|LTS?|L)\b", nombre))
    if pesos:
        m = pesos[-1]
        partes.insert(0, peso(m.group(1), m.group(2)))
        nombre = nombre[:m.start()] + nombre[m.end():]

    nombre = re.sub(r"\s*-\s*(UND|UNIDAD|INDIVIDUAL)\b", r" - \1", nombre)
    nombre = re.sub(r"\b(UND|UNIDAD)\b", "", nombre)
    nombre = re.sub(r"\bSIN GLUTEN\b", "", nombre) if "SIN TACC" not in nombre else nombre
    nombre = re.sub(r"(\s*-\s*)+$", "", nombre.strip())
    nombre = re.sub(r"^\s*-\s*", "", nombre)
    nombre = re.sub(r"\s+-\s+-\s+", " - ", nombre)
    nombre = re.sub(r"\(\s*\)", "", nombre)
    nombre = " ".join(nombre.split()).strip(" -")
    if "INDIVIDUAL" in nombre:
        nombre = re.sub(r"\s*-?\s*INDIVIDUAL", "", nombre)
        partes.append("individual")
    if re.search(r"\d+\s+PORCIONES", nombre):
        m = re.search(r"\s*-?\s*(\d+)\s+PORCIONES", nombre)
        partes.append(f"{m.group(1)} porciones")
        nombre = nombre[:m.start()] + nombre[m.end():]
    return " ".join(nombre.split()).strip(" -"), " · ".join(partes), bulto


def categoria_por_nombre(nombre, por_defecto):
    n = nombre
    if re.search(r"HAMBURGUESA NATURAL|VEGGIE|SOJA|GARBANZO|LENTEJ|YAMANI|MIX VEGETALES|ESPINACA, CEBOLLA|0% CARNE|TOFU|FALAFEL|TOQUES|ADUKI|QUINOA|GREEN MOON|ARVEJA|REMOLACHA", n):
        if not re.search(r"MERLUZA|POLLO|SALMON", n):
            return "Veggie"
    if re.search(r"MERLUZA|SALMON|CAMARON|LANGOSTINO|RABAS|CALAMAR|MEJILLON|PULPO|CORNALITO|CHORITO|JIBIA|ALMEJA|TRUCHA|ATUN|VIEIRA|SURIMI|KANIKAMA|ABADEJO|LENGUADO|PESCADO|TRONKITO|PAELLA|FRUTOS DE MAR|CALAMARETE|TENTACULO|BOTONES DE JIBIA", n):
        return "Pescados y mariscos"
    if "POLLO" in n and not re.search(r"CROQUETA|SORRENTINO|TARTA|CANELON", n):
        return "Pollo"
    if re.search(r"CROQUETA|TORTILLA|ARROLLADITO|BASTONES DE MUZZARELLA|BARRITAS DE MUZZARELLA", n):
        return "Croquetas y bocaditos"
    if re.search(r"PAPA|BATATA", n) and "TORTILLA" not in n:
        return "Papas"
    if re.search(r"CHIPA|CHURRO|MEDIALUNA|CRIOLLO|\bPAN\b", n):
        return "Panificados"
    return por_defecto


def fila(canal, venta, categoria, marca, producto, presentacion="", unidad="unidad",
         kilos_caja="", precio_nuevo=None, precio_fijo=None, stock=True, codigo=""):
    if precio_nuevo and not precio_fijo:
        precio_fijo = precio_nuevo * (1 - DESCUENTO_FIJO_POR_DEFECTO)
    if precio_nuevo and precio_fijo and precio_fijo > precio_nuevo:
        precio_fijo = precio_nuevo * (1 - DESCUENTO_FIJO_POR_DEFECTO)
    return {
        "canal": canal, "venta": venta, "categoria": categoria, "marca": marca,
        "producto": producto, "presentacion": presentacion, "unidad": unidad,
        "kilos_caja": kilos_caja,
        "precio_nuevo": round(precio_nuevo) if precio_nuevo else "",
        "precio_fijo": round(precio_fijo) if precio_fijo else "",
        "stock": "si" if stock else "no", "imagen": "", "codigo": codigo,
    }


def producto_envasado(nombre_crudo, codigo, categoria_hoja, marca_hoja=""):
    nombre = corregir(nombre_crudo)
    marca = detectar_marca(nombre, codigo, marca_hoja)
    nombre, presentacion, bulto = separar_presentacion(quitar_marca(nombre))
    categoria = categoria_por_nombre(nombre, categoria_hoja)
    return marca, titulo(nombre), presentacion, bulto, categoria


# ---------- CATALOGO NUEVO: minorista, empaquetado ----------

def leer_minorista(libro):
    filas = []
    # hoja: (col nombre, col precio, col descuento, categoría por defecto, marca)
    hojas = {
        "REBOZADO ": (2, 4, 5, "Croquetas y bocaditos", ""),
        "PESCADOS": (2, 4, 5, "Pescados y mariscos", ""),
        "FRUTAS Y POSTRES": (2, 3, 5, "Frutas", "Karinat"),
        "PAPAS": (2, 4, None, "Papas", "Finca Balcarce"),
        "PASTAS": (2, 4, None, "Pastas", ""),
        "HAMBURGUESA VEGGIE": (2, 3, 5, "Veggie", ""),
        "SIN GLUTEN": (3, 4, None, "Sin gluten", "Kubi"),
        "TARTAS- VEGGIE": (3, 4, None, "Tartas", "All Frozen"),
    }
    secciones_frutas = {"FRUTAS": "Frutas", "FRUTA GASTRONOMICA": "Frutas",
                        "HELADOS": "Helados y postres", "POSTRES": "Helados y postres"}
    for hoja, (c_nom, c_pre, c_desc, cat_hoja, marca_hoja) in hojas.items():
        ws = libro[hoja]
        categoria = cat_hoja
        for celdas in ws.iter_rows(values_only=True):
            celdas = list(celdas) + [None] * 10
            nombre = texto(celdas[c_nom - 1])
            if hoja == "PASTAS" and not nombre and texto(celdas[0]) == "PASTAS":
                continue
            if hoja == "PASTAS":
                nombre = texto(celdas[1])
            if nombre in secciones_frutas:
                categoria = secciones_frutas[nombre]
                continue
            if not nombre or nombre.upper() in ("REBOZADO", "PESCADO", "PRODUCTO", "PAPAS", "TARTAS", "FRUTAS"):
                continue
            precio = numero(celdas[c_pre - 1])
            if hoja == "SIN GLUTEN" and not precio:
                continue  # filas sin precio cargado (#DIV/0! en el Excel)
            descuento = numero(celdas[c_desc - 1]) if c_desc else None
            fijo = precio - descuento if precio and descuento else None
            codigo = texto(celdas[0]) if hoja != "SIN GLUTEN" else ""
            if codigo.upper() in ("CODIGO", "PASTAS") or len(codigo) < 6:
                codigo = ""
            cat_base = categoria if hoja == "FRUTAS Y POSTRES" else cat_hoja
            marca, producto, presentacion, _, cat = producto_envasado(nombre, codigo, cat_base, marca_hoja)
            if hoja in ("FRUTAS Y POSTRES", "SIN GLUTEN", "PASTAS", "TARTAS- VEGGIE", "HAMBURGUESA VEGGIE"):
                cat = cat_base if hoja != "HAMBURGUESA VEGGIE" else "Veggie"
            filas.append(fila("minorista", "empaquetado", cat, marca, producto, presentacion,
                              precio_nuevo=precio, precio_fijo=fijo, codigo=codigo))
    return filas


def leer_minorista_extra(libro):
    """Productos de 'FRUTA POR MAYOR' que no aparecen en el Excel de granel."""
    filas = []
    ws = libro["FRUTA POR MAYOR"]
    for nombre, precio in ws.iter_rows(values_only=True):
        nombre = texto(nombre)
        if nombre in ("FINCA BALCARCE - PAPA BASTON  2.5K (4)", "CHIPA SIN GLUTEN - KUBI - 10U",
                      "EMPANADAS DE MATAMBRE SIN GLUTEN - KUBI"):
            marca, producto, presentacion, bulto, cat = producto_envasado(nombre, "", "Sin gluten")
            if "CHIPA" in nombre or "EMPANADAS" in nombre:
                cat = "Sin gluten"
            filas.append(fila("mayorista", "empaquetado", cat, marca, producto,
                              presentacion + (f" · caja x{bulto}" if bulto else ""), precio_nuevo=numero(precio)))
    return filas


# ---------- CATALOGO A GRANEL: mayorista ----------

def leer_artico(ws):
    filas = []
    seccion = ""
    for kilos, nombre, caja, kilo, *_ in ws.iter_rows(min_row=2, values_only=True):
        nombre = texto(nombre)
        if not nombre:
            continue
        if kilos is None and caja is None:
            seccion = nombre
            continue
        nombre = corregir(nombre)
        detalle = ""
        m = re.search(r"\(([^)]*)\)\s*$", nombre)
        if m and not re.fullmatch(r"\d+", m.group(1)):
            detalle = m.group(1)
            nombre = nombre[:m.start()]
        detalle = detalle.replace(" X ", " x ").replace("x ", "x ").lower()
        detalle = re.sub(r"(\d)\s*kg", r"\1 kg", detalle)
        detalle = re.sub(r"(\d)\s*grs?\b", r"\1 g", detalle)
        detalle = re.sub(r"\bu\b|unidades?\b", "u.", detalle).replace("paq.", "paquete")
        stock = not (isinstance(caja, str) and "STOCK" in caja.upper()) and not (caja is None and kilo is None)
        al_peso = isinstance(kilos, str)
        precio_kg = numero(kilo)
        if precio_kg is None and numero(caja) and numero(kilos):
            precio_kg = caja / kilos
        presentacion = "Al peso" if al_peso else f"Caja {str(kilos).replace('.', ',')} kg"
        if detalle:
            presentacion += f" · {detalle}"
        de_mar = re.search(r"MAR|SALMON|MERLUZA|CALAMAR|MEJILLON|PULPO|CAMARON|JIBIA|IMPORTADOS|NATURAL", seccion)
        categoria = categoria_por_nombre(nombre, "Pescados y mariscos" if de_mar else "Croquetas y bocaditos")
        marca = detectar_marca(nombre, "", "Ártico")
        producto = titulo(quitar_marca(" ".join(nombre.split())).strip(" -."))
        filas.append(fila("mayorista", "granel", categoria, marca, producto, presentacion, unidad="kg",
                          kilos_caja="" if al_peso else kilos, precio_nuevo=precio_kg if stock else None,
                          stock=stock))
    return filas


def leer_superbe(ws):
    filas = []
    for nombre, presentacion, precio, *_ in ws.iter_rows(min_row=2, values_only=True):
        nombre, presentacion = texto(nombre), texto(presentacion)
        if not nombre or not presentacion:
            continue
        nombre = corregir(nombre).replace(":", "")
        stock = not (isinstance(precio, str) and "STOCK" in precio.upper())
        kilos = re.search(r"(\d+(?:[.,]\d+)?)\s*KG", presentacion.upper())
        kilos = float(kilos.group(1).replace(",", ".")) if kilos else ""
        if kilos and kilos == int(kilos):
            kilos = int(kilos)
        detalle = re.sub(r"CAJA\s*\*\s*(\d+)U", r"\1 u.", presentacion.upper()).replace(" - ", " · ")
        detalle = re.sub(r"(\d)\s*KG", r"\1 kg", detalle)
        peso_txt = ""
        m = re.search(r"\s*X\s?(\d+)\s*GRS?\b", nombre)
        if m:
            peso_txt = f"{m.group(1)} g"
            nombre = nombre[:m.start()] + nombre[m.end():]
        producto = " ".join(nombre.split())
        texto_pres = f"Caja {detalle}" if not detalle.startswith(tuple("0123456789")) or "u." in detalle else f"Caja {detalle}"
        if peso_txt:
            texto_pres = f"{peso_txt} · {texto_pres}"
        filas.append(fila("mayorista", "granel", "Pescados y mariscos", "Superbe", titulo(producto), texto_pres,
                          unidad="kg", kilos_caja=kilos, precio_nuevo=numero(precio), stock=stock))
    return filas


def leer_casa_vegana(ws):
    filas = []
    for u_blister, u_granel, nombre, blister_caja, granel_caja, *_ in ws.iter_rows(min_row=3, values_only=True):
        nombre = texto(nombre)
        if not nombre:
            continue
        producto = titulo(quitar_marca(corregir(nombre)).replace(" 4U", ""))
        if numero(blister_caja):
            filas.append(fila("mayorista", "empaquetado", "Veggie", "Casa Vegana", producto,
                              f"Caja x{u_blister} u. en blíster", unidad="caja", precio_nuevo=blister_caja))
        if numero(granel_caja):
            filas.append(fila("mayorista", "granel", "Veggie", "Casa Vegana", producto,
                              f"Caja x{u_granel} u. a granel", unidad="caja", precio_nuevo=granel_caja))
    return filas


def leer_framfrites(ws):
    filas = []
    seccion = ""
    for nombre, precio in ws.iter_rows(values_only=True):
        nombre = texto(nombre)
        if not nombre:
            continue
        if precio is None or nombre == "EXPRIMIDOS":
            seccion = nombre
            continue
        nombre = corregir(nombre)
        n = nombre
        presentacion = ""
        m = re.search(r"\s+X\s+(\d+)\s*U\.?\s*$", n)
        if m:
            presentacion = f"Caja x{m.group(1)} u."
            n = n[:m.start()]
        m = re.search(r"\s+(\d+)\s*U\.\s*X\s*([\d.,]+)\s*LTS\.?", n)
        if m:
            presentacion = f"{m.group(1)} u. x {m.group(2)} L"
            n = n[:m.start()]
        m = re.search(r"\s+X\s+(\d+)\s*BL\.?\s*X\s*([\d.,]+)\s*K(?:G)?", n)
        if m:
            presentacion = f"Caja {m.group(1)} bolsas x {m.group(2).replace('.', ',')} kg"
            n = n[:m.start()]
        m = re.search(r"\s+(\d+)\s*X\s*([\d.,]+)\s*K$", n)
        if m:
            presentacion = f"Caja {m.group(1)} bolsas x {m.group(2).replace('.', ',')} kg"
            n = n[:m.start()]
        m = re.search(r"\s+X\s+(\d+)\s*KG\.?\s*GRANEL", n)
        if m:
            presentacion = f"Bolsa {m.group(1)} kg"
            n = n[:m.start()]
        n = n.replace("MED.", "MEDIALUNA").replace("UNR", "").replace("FARMFRITES", "").replace("FARM FRITES", "")
        es_papa = nombre.startswith("PAPA")
        marca = "Farm Frites" if es_papa else ""
        producto = titulo(" ".join(n.split()))
        if es_papa:
            producto = "Papas " + producto.replace("Papa ", "").replace("Star Frites", "Star Frites").strip()
        venta = "empaquetado" if seccion == "EXPRIMIDOS" else "granel"
        categoria = "Bebidas" if seccion == "EXPRIMIDOS" else ("Papas" if es_papa else "Panificados")
        filas.append(fila("mayorista", venta, categoria, marca, producto, presentacion,
                          unidad="caja", precio_nuevo=numero(precio)))
    return filas


def leer_karinat(ws):
    filas = []
    categoria = "Frutas"
    secciones = {"FRUTAS": "Frutas", "FRUTA GASTRONOMICA": "Frutas", "HELADOS": "Helados y postres", "POSTRES": "Helados y postres"}
    for nombre, precio in ws.iter_rows(values_only=True):
        nombre = texto(nombre)
        if nombre in secciones:
            categoria = secciones[nombre]
            continue
        if not nombre:
            continue
        marca, producto, presentacion, bulto, _ = producto_envasado(nombre, "", categoria, "Karinat")
        if bulto:
            presentacion += f" · caja x{bulto}"
        filas.append(fila("mayorista", "empaquetado", categoria, marca or "Karinat", producto,
                          presentacion.strip(" ·"), precio_nuevo=numero(precio)))
    return filas


def leer_sin_gluten(ws):
    filas = []
    for _, nombre, precio in ws.iter_rows(min_row=2, values_only=True):
        nombre = texto(nombre)
        if not nombre or nombre == "SIN GLUTEN":
            continue
        marca, producto, presentacion, _, _ = producto_envasado(nombre, "", "Sin gluten", "Kubi")
        filas.append(fila("mayorista", "empaquetado", "Sin gluten", marca, producto, presentacion,
                          precio_nuevo=numero(precio)))
    return filas


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    nuevo = openpyxl.load_workbook(sys.argv[1], data_only=True)
    granel = openpyxl.load_workbook(sys.argv[2], data_only=True)

    filas = leer_minorista(nuevo)
    filas += leer_karinat(granel["KARINAT"])
    filas += leer_sin_gluten(granel["SIN GLUTEN"])
    filas += leer_minorista_extra(nuevo)
    filas += leer_casa_vegana(granel["CASA VEGANA"])
    filas += leer_framfrites(granel["FRAMFRITES"])
    filas += leer_artico(granel["ARTICO"])
    filas += leer_superbe(granel["SUPERBE"])

    filas.sort(key=lambda f: (f["canal"] != "minorista", f["venta"] != "empaquetado",
                              ORDEN_CATEGORIAS.index(f["categoria"]) if f["categoria"] in ORDEN_CATEGORIAS else 99))

    productos = []
    usados = set()
    for f in filas:
        base = re.sub(r"[^a-z0-9]+", "-", f"{f['canal'][:3]}-{f['venta'][:3]}-{f['producto']}".lower()
                      .translate(str.maketrans("áéíóúüñ", "aeiouun"))).strip("-")
        id_producto, n = base, 2
        while id_producto in usados:
            id_producto, n = f"{base}-{n}", n + 1
        usados.add(id_producto)
        productos.append({
            "id": id_producto,
            "canal": f["canal"], "venta": f["venta"], "categoria": f["categoria"],
            "marca": f["marca"], "producto": f["producto"], "presentacion": f["presentacion"],
            "unidad": f["unidad"],
            "kilos_caja": f["kilos_caja"] if f["kilos_caja"] != "" else None,
            "precio_nuevo": f["precio_nuevo"] or None,
            "precio_fijo": f["precio_fijo"] or None,
            "stock": f["stock"] == "si",
            "imagen": "",
        })

    destino = Path(__file__).resolve().parent.parent / "datos" / "productos.json"
    destino.parent.mkdir(exist_ok=True)
    with open(destino, "w", encoding="utf-8") as archivo:
        json.dump({"actualizado": "", "productos": productos}, archivo, ensure_ascii=False, indent=1)
    print(f"{len(productos)} productos -> {destino}")

if __name__ == "__main__":
    main()
